# -*- coding: utf-8 -*-
# Supplementary Methods S1 재배치: (버그) 그림 뒤 + 제목·본문 역순 → 첫 그림 문단 앞, [제목, 본문] 순.
import copy
from docx import Document
from docx.text.paragraph import Paragraph

S1_HEAD = "Supplementary Methods S1. Sensor-to-neighbourhood matching"
BODY_MARK = "Sensor metadata contain street addresses but no administrative-neighbourhood name. Each sensor was therefore"

SUPPS = [
    r"C:\Users\wh850\Research\assets\31_환경소음_이동성\01_논문작업\Supplementary_20260726_231342.docx",
    r"C:\Users\wh850\Research\assets\31_환경소음_이동성\260806_LUP_투고\04_Supplementary.docx",
]
for path in SUPPS:
    doc = Document(path)
    head = [p for p in doc.paragraphs if p.text.strip() == S1_HEAD]
    body = [p for p in doc.paragraphs if p.text.strip().startswith(BODY_MARK)]
    assert len(head) == 1 and len(body) == 1, f"{path}: head {len(head)} body {len(body)}"
    body_text = body[0].text
    # 캡션 문단(서식 원본) 확보 후 기존 삽입분 제거
    cap = [p for p in doc.paragraphs if p.text.strip().startswith("Supplementary Fig. S1")]
    assert len(cap) == 1
    for p in (head[0], body[0]):
        p._element.getparent().remove(p._element)
    # 첫 그림(drawing) 문단 앞에 [제목, 본문] 순 삽입
    draw = None
    for p in doc.paragraphs:
        if p._p.xpath('.//w:drawing'):
            draw = p; break
    assert draw is not None, "drawing 문단 미발견"
    for txt, bold in ((S1_HEAD, True), (body_text, False)):
        new_p = copy.deepcopy(cap[0]._p)
        draw._p.addprevious(new_p)
        np = Paragraph(new_p, draw._parent)
        for i, r in enumerate(np.runs):
            r.text = txt if i == 0 else ""
        for r in np.runs:
            r.font.bold = bold
            r.font.italic = False
    doc.save(path)
    print("재배치 OK:", path.split(chr(92))[-1])
