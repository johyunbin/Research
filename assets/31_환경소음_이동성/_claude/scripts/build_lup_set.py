# -*- coding: utf-8 -*-
# LUP(Landscape and Urban Planning) 재투고 세트 구성 — 260802 SCS 제출 세트 기반.
# 00 커버레터: 수신 3인 Co-EiC(2026-08-06 라이브 실측)·저널명·날짜·fit 문단 교체(서식 보존).
# 03 원고: 키워드 "smart-city monitoring"→"land use" (분량 감축은 별도 결정).
# 01·02·04·figures: 원본 복사.
import os, shutil
from docx import Document

SRC = r"C:\Users\wh850\Research\assets\31_환경소음_이동성\260802_SCS_1차투고"
DST = r"C:\Users\wh850\Research\assets\31_환경소음_이동성\260806_LUP_투고"
os.makedirs(DST, exist_ok=True)

for f in ["01_Highlights.docx", "02_Title Page.docx", "04_Supplementary.docx"]:
    shutil.copy2(os.path.join(SRC, f), os.path.join(DST, f))
if not os.path.isdir(os.path.join(DST, "figures")):
    shutil.copytree(os.path.join(SRC, "figures"), os.path.join(DST, "figures"))
print("01/02/04/figures 복사")

# ---- 00 커버레터 ----
cl = os.path.join(DST, "00_Cover_Letter.docx")
shutil.copy2(os.path.join(SRC, "00_Cover_Letter.docx"), cl)
doc = Document(cl)

def set_text(p, new):
    for i, r in enumerate(p.runs):
        r.text = new if i == 0 else ""

FIT_NEW = ("We believe this work is well suited to Landscape and Urban Planning. The sound environment is "
           "part of the everyday quality of urban landscapes, and the paper examines how it responds when "
           "the distribution of human activity across neighbourhoods shifts, at the scale planning works "
           "with: the neighbourhood and its land use. The broadly uniform response across commercial, mixed "
           "and residential areas bears directly on the choice between land-use-differentiated and citywide "
           "noise measures. The estimate gives planners a quantitative basis for judging the noise "
           "co-benefits of mobility-oriented interventions, and the drift-aware design offers a "
           "transferable template for using city-scale sensor networks in planning research. All raw data "
           "are openly available from public portals.")

n = 0
zhou_p = None
for p in doc.paragraphs:
    t = p.text.strip()
    if t == "August 2, 2026":
        set_text(p, "August 6, 2026"); n += 1
    elif t == "Prof. Enrico Fabrizio":
        set_text(p, "Prof. Christian Albert"); n += 1
    elif t == "Prof. Fariborz Haghighat":
        set_text(p, "Prof. Iryna Dronova"); n += 1
    elif t == "Prof. Ryozo Ooka":
        set_text(p, "Prof. Weiqi Zhou"); zhou_p = p; n += 1
    elif t == "Sustainable Cities and Society":
        set_text(p, "Landscape and Urban Planning"); n += 1
    elif "for consideration as an original research article in" in p.text:
        # 제목 볼드·저널 이탤릭 run 서식 보존: run 단위 치환
        for r in p.runs:
            if "original research article" in r.text:
                r.text = r.text.replace("as an original research article in", "as a Research Paper in")
            if "Sustainable Cities and Society" in r.text:
                r.text = r.text.replace("Sustainable Cities and Society", "Landscape and Urban Planning")
        n += 1
    elif "much of sustainable urban policy" in p.text:
        for r in p.runs:
            if "much of sustainable urban policy" in r.text:
                r.text = r.text.replace("much of sustainable urban policy", "much of urban planning and policy")
        n += 1
    elif p.text.startswith("We believe this work is well suited to Sustainable Cities and Society."):
        sz = None
        for r in p.runs:
            if r.font.size is not None:
                sz = r.font.size.pt; break
        for i, r in enumerate(p.runs):
            r.text = FIT_NEW if i == 0 else ""
        n += 1
assert n == 8, f"커버레터 치환 {n}건 (기대 8)"
assert zhou_p is not None

# Zhou 아래에 Co-Editors-in-Chief 줄 삽입 (Zhou 문단 복제)
import copy
new_p = copy.deepcopy(zhou_p._p)
zhou_p._p.addnext(new_p)
from docx.text.paragraph import Paragraph
np = Paragraph(new_p, zhou_p._parent)
set_text(np, "Co-Editors-in-Chief")

doc.save(cl)
print("00 커버레터 LUP판 저장")

# ---- 03 원고: 키워드 교체 ----
ms = os.path.join(DST, "03_Manuscript.docx")
shutil.copy2(os.path.join(SRC, "03_Manuscript.docx"), ms)
doc = Document(ms)
n = 0
for p in doc.paragraphs:
    if p.text.startswith("Keywords:"):
        for r in p.runs:
            if "smart-city monitoring" in r.text:
                r.text = r.text.replace("smart-city monitoring", "land use"); n += 1
        break
assert n == 1, f"키워드 치환 {n}건"
doc.save(ms)
print("03 키워드 교체 저장")
