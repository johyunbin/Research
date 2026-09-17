# -*- coding: utf-8 -*-
"""
Paper32 — 사용자가 추가로 내려받은 PDF를 미확보 목록과 대조해 수집논문_PDF 라이브러리로 편입 (2026-09-18)

대조 풀(전문 미확보 레코드 전부):
  DB  = fulltext/db_oa_retrieval_20260917_233746.csv (89건; OK_XML 11건 포함 — PDF를 받았으면 PDF로 교체)
  CT  = fulltext/ct_retrieval_final.csv state=not-retrieved (25건)
  OAS = fulltext/oa_supp_retrieval.csv result in {no-oa-pdf, http403, http418} (8건)
원본 파일: 01_논문작업/*.pdf 중 2026-09-17 23:50 KST 이후 저장분(LUP 비교 논문 4편 제외) + _claude/db_oa_pdf/*.pdf(자동 확보 10건)
대조: 앞 2쪽 본문의 DOI 일치 + 제목 핵심어 일치율. 둘 중 하나로 유일하게 확정될 때만 편입.
명명: 년도_저널축약_저자_제목.pdf (consolidate_ct_into_library.py 규칙 재사용, 메타데이터는 Crossref)
동작: 사용자 파일은 라이브러리로 이동, 자동 확보분은 복사. 대조 실패·중복은 제자리에 두고 보고.
출력: fulltext/manual_pdf_ingest_<타임코드>.csv · 01_논문작업/미확보_원문_목록_<타임코드>.xlsx
"""
import csv, html, json, os, re, shutil, sys, time, unicodedata, urllib.parse, urllib.request, hashlib
from datetime import datetime, timezone, timedelta

import fitz
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

sys.stdout.reconfigure(encoding="utf-8")
BASE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(BASE)
FT = os.path.join(BASE, "fulltext")
LIB = os.path.join(ROOT, "수집논문_PDF")
WORK = os.path.join(ROOT, "01_논문작업")
KST = timezone(timedelta(hours=9))
STAMP = datetime.now(KST).strftime("%Y%m%d_%H%M%S")
CUTOFF = datetime(2026, 9, 17, 23, 50, tzinfo=KST).timestamp()
LUP_COMPARISON = {"1-s2.0-S0169204624000707-main.pdf", "1-s2.0-S0169204625001148-main.pdf",
                  "1-s2.0-S0169204625002051-main.pdf", "1-s2.0-S0169204626001313-main.pdf"}
UA = {"User-Agent": "paper32-systematic-review/1.0"}
sys.path.insert(0, BASE)
from consolidate_ct_into_library import abbrev, safe  # noqa: E402  같은 명명 규칙 재사용

STOP = {"with", "from", "that", "this", "their", "into", "between", "among", "based", "study", "effects",
        "effect", "urban", "case", "using", "analysis"}


def norm_doi(d):
    return (d or "").strip().lower().rstrip(".")


def words(t):
    t = unicodedata.normalize("NFKC", t or "").lower()
    return {w for w in re.findall(r"[a-z]{4,}", t) if w not in STOP}


def load_pool():
    pool = []
    for r in csv.DictReader(open(os.path.join(FT, "db_oa_retrieval_20260917_233746.csv"), encoding="utf-8-sig")):
        pool.append({"branch": "DB", "id": r["no"], "prio": r["verdict"], "year": r["year"], "journal": r["journal"],
                     "title": r["title"], "doi": r["doi"], "oa": r["oa_status"], "auto": r["status"],
                     "note": "8월 미확보 · 9/17 무료 원문 자동 시도 실패"})
    for r in csv.DictReader(open(os.path.join(FT, "ct_retrieval_final.csv"), encoding="utf-8-sig")):
        if r["state"] == "not-retrieved":
            pool.append({"branch": "CT", "id": f"CT{int(r['rec']):04d}", "prio": r["screen"], "year": r["year"],
                         "journal": r["journal"], "title": r["title"], "doi": r["doi"], "oa": r["oa_status"],
                         "auto": "NOT_RETRIEVED", "note": "8월 미확보: " + (r["reason"] or "사유 미기록")})
    oas_reason = {"no-oa-pdf": "무료 원문 없음", "http403": "출판사 접근 차단(403)", "http418": "출판사 봇 차단(418)"}
    for r in csv.DictReader(open(os.path.join(FT, "oa_supp_retrieval.csv"), encoding="utf-8-sig")):
        if r["result"] in oas_reason:
            pool.append({"branch": "OAS", "id": f"OAS{int(r['sid']):04d}", "prio": r["verdict"], "year": r["year"],
                         "journal": r["journal"], "title": r["title"], "doi": r["doi"], "oa": r["oa_status"],
                         "auto": "NOT_RETRIEVED", "note": "8월 미확보: " + oas_reason[r["result"]]})
    return pool


def completeness_audit(have, xml_only, missing, lib_titles):
    """경로별 전문 확보 대상 전체 = 8월 확보 + 이번 추가 확보 + XML 확보 + 미확보 (+ 보조검색 사전 배제) 인지 확인.
    8월 확보분이 라이브러리에 실제로 있는지도 파일명으로 대조한다."""
    def in_lib(title):
        key = compact(safe(title, 70))[:40]
        return bool(key) and any(t.startswith(key) for t in lib_titles)

    rows = []
    db = list(csv.DictReader(open(os.path.join(FT, "fulltext_status_20260803.csv"), encoding="utf-8-sig")))
    ct = list(csv.DictReader(open(os.path.join(FT, "ct_retrieval_final.csv"), encoding="utf-8-sig")))
    oas = list(csv.DictReader(open(os.path.join(FT, "oa_supp_retrieval.csv"), encoding="utf-8-sig")))
    oas_dir = os.path.join(BASE, "oa_supp_pdf")
    oas_files = {f[:7] for f in os.listdir(oas_dir)} if os.path.isdir(oas_dir) else set()
    for name, recs, orig_ok, rid, pre_ex, lib_check in [
        ("데이터베이스 검색", db, lambda r: r["pdf"] == "Y", lambda r: r["no"], lambda r: False,
         lambda r: in_lib(r["title"])),
        ("인용 추적", ct, lambda r: r["state"] == "retrieved", lambda r: f"CT{int(r['rec']):04d}", lambda r: False,
         lambda r: in_lib(r["title"])),
        ("보조 검색", oas, lambda r: r["result"] == "ok", lambda r: f"OAS{int(r['sid']):04d}",
         lambda r: r["result"] == "prescreen-exclude", lambda r: rid_oas(r) in oas_files or in_lib(r["title"])),
    ]:
        orig = [r for r in recs if orig_ok(r)]
        pre = [r for r in recs if pre_ex(r)]
        rest = [r for r in recs if not orig_ok(r) and not pre_ex(r)]
        added = [r for r in rest if rid(r) in have]
        xml = [r for r in rest if rid(r) in {p["id"] for p in xml_only}]
        miss = [r for r in rest if rid(r) in {p["id"] for p in missing}]
        unaccounted = len(rest) - len(added) - len(xml) - len(miss)
        orig_found = sum(1 for r in orig if lib_check(r))
        rows.append({"경로": name, "전문 확보 대상": len(recs), "8월 확보": len(orig), "8월 확보 중 파일 확인": orig_found,
                     "이번 추가 확보(PDF)": len(added), "XML 확보": len(xml), "사전 배제(학술지 아님)": len(pre),
                     "미확보": len(miss), "누락(0이어야 함)": unaccounted})
    return rows


def rid_oas(r):
    return f"OAS{int(r['sid']):04d}"


def pdf_text(path):
    doc = fitz.open(path)
    pages = [unicodedata.normalize("NFKC", doc[i].get_text()) for i in range(min(2, doc.page_count))]  # ﬁ 합자 풀기
    pages = [re.sub(r"-\s*\n\s*", "", p) for p in pages]
    meta_title = unicodedata.normalize("NFKC", (doc.metadata or {}).get("title") or "")
    return " ".join(pages), (pages[0] if pages else "")[:2500], doc.page_count, meta_title


def match(path, pool):
    """DOI 일치(앞 2쪽 또는 파일명)가 있으면 우선. 제목 대조는 첫 쪽 머리(2,500자)에서만 한다 —
    본문·참고문헌에 인용된 다른 논문 제목과 섞이지 않게(예: 같은 저자의 전년도 논문)."""
    text, head, pages, meta_title = pdf_text(path)
    low, head_low = text.lower(), (head + " " + meta_title).lower()
    dois = {norm_doi(d) for d in re.findall(r"10\.\d{4,9}/[^\s\"<>)\]]+", low)}
    fname = os.path.basename(path).lower()
    scored = []
    for p in pool:
        tw = words(p["title"])
        if not tw:
            continue
        hit_all = sum(1 for w in tw if w in low) / len(tw)
        hit_head = sum(1 for w in tw if w in head_low) / len(tw)
        d = norm_doi(p["doi"])
        doi_hit = bool(d) and (d in dois or d.replace("/", "@") in fname or d.split("/")[-1] in fname)
        exact = compact(p["title"])[:60] in compact(head + " " + meta_title)   # 제목 앞 60자가 연속으로 나타나는가
        scored.append((doi_hit and hit_all >= 0.4, exact, hit_head, p))
    strong = [s for s in scored if s[0] or s[1] or s[2] >= 0.75]
    strong.sort(key=lambda s: (s[0], s[1], s[2]), reverse=True)
    if not strong:
        best = max(scored, key=lambda s: (s[0], s[1], s[2]))
        return None, f"no_match(best {best[3]['id']} doi={best[0]} title={best[2]:.0%})", pages
    if len(strong) > 1 and strong[0][:2] == strong[1][:2] and strong[1][2] >= strong[0][2] - 0.05:
        return None, f"ambiguous({strong[0][3]['id']},{strong[1][3]['id']})", pages
    return strong[0][3], f"doi={strong[0][0]} exact={strong[0][1]} title_head={strong[0][2]:.0%}", pages


def compact(s):
    return re.sub(r"[^a-z0-9]+", "", unicodedata.normalize("NFKC", s or "").lower())


def crossref_meta(doi, title):
    try:
        if doi:
            url = "https://api.crossref.org/works/" + urllib.parse.quote(doi)
            d = json.loads(urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=30).read())["message"]
        else:
            url = ("https://api.crossref.org/works?rows=1&query.bibliographic=" + urllib.parse.quote(title))
            items = json.loads(urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=30).read())["message"]["items"]
            if not items:
                return None
            d = items[0]
            if len(words(title) & words((d.get("title") or [""])[0])) / max(1, len(words(title))) < 0.8:
                return None
    except Exception:
        return None
    finally:
        time.sleep(0.4)
    auths = d.get("author") or []
    fam = lambda a: a.get("family") or a.get("name") or ""
    if not auths:
        astr = "Anon"
    elif len(auths) == 1:
        astr = fam(auths[0])
    elif len(auths) == 2:
        astr = f"{fam(auths[0])} & {fam(auths[1])}"
    else:
        astr = f"{fam(auths[0])} et al"
    return {"journal": (d.get("container-title") or [""])[0], "authors": astr, "title": (d.get("title") or [""])[0]}


def sha(path):
    return hashlib.sha256(open(path, "rb").read()).hexdigest()


def main():
    pool = load_pool()
    lib_hashes = {sha(os.path.join(LIB, f)): f for f in os.listdir(LIB) if f.lower().endswith(".pdf")}
    sources = []
    for f in sorted(os.listdir(WORK)):
        p = os.path.join(WORK, f)
        if f.lower().endswith(".pdf") and f not in LUP_COMPARISON and os.path.getmtime(p) >= CUTOFF:
            sources.append(("user", p))
    auto_dir = os.path.join(BASE, "db_oa_pdf")
    for f in sorted(os.listdir(auto_dir)):
        if f.lower().endswith(".pdf"):
            sources.append(("auto", os.path.join(auto_dir, f)))
    print(f"대조 풀 {len(pool)}건 · 원본 PDF {len(sources)}개(사용자 {sum(s[0]=='user' for s in sources)})")

    # 자동 대조가 못 잡는 경우의 수기 지정(내용 확인 후): 학회 초록집 2쪽 중간에 대상 초록이 있는 파일
    MANUAL = {"haddak2017.pdf": "OAS0222"}
    by_id = {p["id"]: p for p in pool}
    log, got = [], {}
    for origin, path in sources:
        if os.path.basename(path) in MANUAL:
            rec, pages = by_id[MANUAL[os.path.basename(path)]], fitz.open(path).page_count
            text = unicodedata.normalize("NFKC", " ".join(pg.get_text() for pg in fitz.open(path))).lower()
            assert compact(rec["title"])[:50] in compact(text), "수기 지정 파일에 대상 제목이 없다"
            why = "manual(title found in text)"
        else:
            rec, why, pages = match(path, pool)
        row = {"origin": origin, "file": os.path.basename(path), "pages": pages, "match": why,
               "branch": "", "id": "", "doi": "", "action": "", "library_name": ""}
        if rec is None:
            row["action"] = "left_in_place"
            log.append(row)
            print(f"  ✗ {row['file'][:50]:50} {why}")
            continue
        row.update(branch=rec["branch"], id=rec["id"], doi=rec["doi"])
        h = sha(path)
        if rec["id"] in got:
            row["action"] = "duplicate_of_" + got[rec["id"]] + ("(identical)" if got.get(rec["id"] + "#h") == h else "")
            log.append(row)
            print(f"  = {row['file'][:50]:50} 중복 → {rec['id']}")
            continue
        if h in lib_hashes:
            row.update(action="already_in_library", library_name=lib_hashes[h])
            got[rec["id"]], got[rec["id"] + "#h"] = row["file"], h
            log.append(row)
            continue
        meta = crossref_meta(rec["doi"].strip(), rec["title"]) or {}
        name = (f"{rec['year']}_{abbrev(meta.get('journal') or rec['journal'])}_{safe(meta.get('authors') or 'Unknown', 40)}"
                f"_{safe(meta.get('title') or rec['title'], 70)}.pdf")
        dst = os.path.join(LIB, name)
        if os.path.exists(dst):
            row.update(action="name_exists", library_name=name)
        elif origin == "user":
            shutil.move(path, dst)
            row.update(action="moved", library_name=name)
        else:
            shutil.copy2(path, dst)
            row.update(action="copied", library_name=name)
        got[rec["id"]], got[rec["id"] + "#h"] = row["file"], h
        log.append(row)
        print(f"  ✓ {row['file'][:40]:40} → [{rec['branch']} {rec['id']}] {name[:70]}")

    with open(os.path.join(FT, f"manual_pdf_ingest_{STAMP}.csv"), "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=list(log[0].keys()))
        w.writeheader()
        w.writerows(log)

    # ── 아직 없는 논문 목록 ──────────────────────────────────────────────
    # 확보 판정 = 지금까지의 모든 편입 기록(이번 실행 포함) ∪ 라이브러리 파일명에 제목 앞부분이 있는 경우
    OK = ("moved", "copied", "already_in_library", "name_exists")
    have = set()
    for f in os.listdir(FT):
        if f.startswith("manual_pdf_ingest_") and f.endswith(".csv"):
            have |= {r["id"] for r in csv.DictReader(open(os.path.join(FT, f), encoding="utf-8-sig"))
                     if r["id"] and r["action"] in OK}
    lib_titles = [compact(re.sub(r"^\d{4}_[^_]*_[^_]*_", "", f[:-4])) for f in os.listdir(LIB) if f.lower().endswith(".pdf")]
    for p in pool:
        key = compact(safe(p["title"], 70))[:40]
        if key and any(t.startswith(key) for t in lib_titles):
            have.add(p["id"])
    missing = [p for p in pool if p["id"] not in have and p["auto"] != "OK_XML"]
    xml_only = [p for p in pool if p["id"] not in have and p["auto"] == "OK_XML"]
    audit = completeness_audit(have, xml_only, missing, lib_titles)
    branch_name = {"DB": "데이터베이스 검색", "CT": "인용 추적(8월 시도)", "OAS": "보조 검색(8월 시도)"}

    def grp(p):
        if p["prio"] in ("INCLUDE", "RETRIEVE"):
            g = (1, "1. 선별 단계 포함 판정")
        elif not p["doi"].strip():
            g = (4, "4. DOI 없음(제목으로 검색)")
        elif p["oa"] in ("closed", ""):
            g = (3, "3. 구독 전용")
        else:
            g = (2, "2. 무료 원문 표시(자동 차단)")
        return g

    missing.sort(key=lambda p: ({"DB": 0, "CT": 1, "OAS": 2}[p["branch"]], grp(p)[0], p["year"]))
    wb = Workbook()
    ws = wb.active
    ws.title = "미확보 원문"
    ws.append(["#", "경로", "구분", "ID", "연도", "학술지", "제목", "DOI", "OA 표시", "이전 시도"])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    for k, p in enumerate(missing, 1):
        doi = p["doi"].strip()
        ws.append([k, branch_name[p["branch"]], grp(p)[1], p["id"], p["year"], html.unescape(html.unescape(p["journal"])), html.unescape(html.unescape(p["title"])), doi or "없음",
                   p["oa"] or "-", p["note"]])
        if doi:
            c = ws.cell(row=k + 1, column=8)
            c.hyperlink = "https://doi.org/" + urllib.parse.quote(doi, safe="/:.;-_")
            c.font = Font(color="0563C1", underline="single")
        for c in ws[k + 1]:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    for col, wdt in zip("ABCDEFGHIJ", [4, 18, 24, 9, 6, 26, 70, 34, 9, 30]):
        ws.column_dimensions[col].width = wdt
    ws.freeze_panes = "D2"
    note = wb.create_sheet("안내")
    lines = [f"아직 전문이 없는 논문 전부 {len(missing)}건입니다 (데이터베이스 {sum(p['branch']=='DB' for p in missing)} · "
             f"인용 추적 {sum(p['branch']=='CT' for p in missing)} · 보조 검색 {sum(p['branch']=='OAS' for p in missing)}).",
             "세 경로에서 전문 확보 대상이었던 레코드를 전부 대조해 만든 목록입니다. 아래 대조표의 '누락'이 모두 0이면 빠진 논문이 없습니다.",
             f"Europe PMC 전문(XML)으로 이미 확보한 {len(xml_only)}건은 평가에 쓸 수 있어 목록에서 뺐습니다.",
             "구분 1: 선별 단계에서 포함(INCLUDE/RETRIEVE)으로 판정 — 가장 먼저 확보",
             "구분 2: OpenAlex 에 무료 원문 표시 — 브라우저로 열면 대부분 받을 수 있음",
             "구분 3: 구독 전용 · 구분 4: DOI 없음",
             "받은 PDF는 01_논문작업 폴더에 넣어 주시면 이름을 맞춰 수집논문_PDF 로 옮깁니다.", "", "경로별 대조표"]
    for i, t in enumerate(lines, 1):
        note.cell(row=i, column=1, value=t)
    r0 = len(lines) + 1
    keys = list(audit[0].keys())
    for j, k in enumerate(keys, 1):
        note.cell(row=r0, column=j, value=k).font = Font(bold=True)
    for i, row in enumerate(audit, 1):
        for j, k in enumerate(keys, 1):
            note.cell(row=r0 + i, column=j, value=row[k])
    note.column_dimensions["A"].width = 18
    for col in "BCDEFGHI":
        note.column_dimensions[col].width = 14
    print("\n경로별 대조표:")
    for row in audit:
        print("  ", row)
    assert all(row["누락(0이어야 함)"] == 0 for row in audit), "대조표에 누락이 있다 — 목록을 내보내기 전에 원인 확인"
    out = os.path.join(WORK, f"미확보_원문_목록_{STAMP}.xlsx")
    wb.save(out)

    from collections import Counter
    print("\n처리:", Counter(r["action"].split("(")[0] for r in log))
    print("확보(경로별):", Counter(r["branch"] for r in log if r["action"] in ("moved", "copied", "already_in_library", "name_exists")))
    print("남은 미확보:", Counter(p["branch"] for p in missing), "· XML만 확보:", len(xml_only))
    print("목록:", out)


if __name__ == "__main__":
    main()
