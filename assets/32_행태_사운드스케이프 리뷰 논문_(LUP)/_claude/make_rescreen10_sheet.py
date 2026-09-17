# -*- coding: utf-8 -*-
"""
Paper32 — 제목·초록 10% 재스크리닝 시트 (OSF 등록본 Screening reliability 이행, 2026-09-17 사용자 결정)

등록본: "a random 10% subsample is re-screened by the same reviewer after an interval to check intra-rater consistency"
- 1,316건에서 무작위 10%(132건)를 고정 시드로 추출한다(재현 가능).
- 시트에는 등록본이 스크리닝 때 보이게 한 필드(제목·초록·학술지·연도·DOI)만 넣고, 이전 판정·사유·방향은 넣지 않는다.
- 판정 비교는 rescreen10_agreement.py 가 원 판정(screening_results_20260802_221808.csv)과 대조해 수행한다.
출력: screening/rescreen10/rescreen10_sheet_<타임코드>.xlsx · rescreen10_sample_<타임코드>.csv
"""
import csv, os, random, sys
from datetime import datetime, timezone, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

sys.stdout.reconfigure(encoding="utf-8")
BASE = os.path.dirname(os.path.abspath(__file__))
POOL = os.path.join(BASE, "screening", "screening_dataset_20260802_221808.csv")
OUTDIR = os.path.join(BASE, "screening", "rescreen10")
SEED = 20260917
FRACTION = 0.10

CRITERIA = [
    ("EX1", "Animal/wildlife subjects (soundscape ecology)", "동물·야생 대상(사운드스케이프 생태학)"),
    ("EX2", "Setting is not an outdoor or semi-outdoor urban/landscape public space (e.g., indoor retail, workplace, clinical, vehicle interior)",
     "옥외·반옥외 도시·경관 공공공간이 아님(실내 상업시설, 직장, 병원, 차량 내부 등)"),
    ("EX3", "No acoustic-environment exposure or manipulation (for reverse-direction records, no behaviour-to-soundscape pathway)",
     "음환경 노출·조작 없음(역방향 레코드는 행태→사운드스케이프 경로 없음)"),
    ("EX4", "No observable behavioural outcome (perception-, annoyance-, preference-, intention-, or physiology-only)",
     "관찰 가능한 행태 결과 없음(지각·성가심·선호·의도·생리 결과만)"),
    ("EX5", "Not an empirical study (review, commentary, pure simulation)", "실증 연구 아님(리뷰, 논평, 순수 시뮬레이션)"),
    ("EX6", "Not a peer-reviewed journal article", "동료심사 학술지 논문 아님"),
    ("EX7", "Not in English", "영어 아님"),
]


def kst_stamp():
    return datetime.now(timezone(timedelta(hours=9))).strftime("%Y%m%d_%H%M%S")


def main():
    pool = list(csv.DictReader(open(POOL, encoding="utf-8-sig")))
    ids = sorted(int(r["no"]) for r in pool)
    by_id = {int(r["no"]): r for r in pool}
    n = round(len(ids) * FRACTION)
    sample = random.Random(SEED).sample(ids, n)

    os.makedirs(OUTDIR, exist_ok=True)
    stamp = kst_stamp()
    with open(os.path.join(OUTDIR, f"rescreen10_sample_{stamp}.csv"), "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["seq", "no", "seed", "pool_size", "sample_size"])
        for i, no in enumerate(sample, 1):
            w.writerow([i, no, SEED, len(ids), n])

    wb = Workbook()
    thin = Side(style="thin", color="BFBFBF")
    head_fill = PatternFill("solid", fgColor="DDEBF7")
    input_fill = PatternFill("solid", fgColor="FFF2CC")

    # ── 안내 시트 ─────────────────────────────────────────────────────────
    ins = wb.active
    ins.title = "Instructions"
    lines = [
        ("제목·초록 10% 재스크리닝 (OSF 등록본 Screening reliability)", True),
        (f"대상: 스크리닝 풀 {len(ids):,}건 중 무작위 {n}건 (시드 {SEED}). 이전 판정은 시트에 넣지 않았습니다.", False),
        ("이전 판정·AI 결과·screening 폴더의 결과 파일을 보지 말고, 제목·초록·학술지·연도만으로 판정해 주세요.", False),
        ("", False),
        ("판정(Decision)", True),
        ("Include = 적격이 분명함(전문 검토로 넘김)", False),
        ("Uncertain = 제목·초록만으로 판단하기 어려움(전문 검토로 넘김, 등록본의 over-inclusive 규칙)", False),
        ("Exclude = 아래 배제 기준 중 하나에 해당. 번호 순서대로 적용하고 처음 해당하는 코드 하나만 적습니다.", False),
        ("", False),
        ("배제 기준 (등록본 'Used exclusion criteria', 번호 순서대로 적용)", True),
    ]
    for code, en, ko in CRITERIA:
        lines.append((f"{code}  {ko}", False))
        lines.append((f"        {en}", False))
    lines += [
        ("", False),
        ("포함 기준 요약 (등록본)", True),
        ("영어 동료심사 학술지 실증 연구 · 옥외/반옥외 도시·경관 공공공간(공원, 가로, 광장, 수변, 캠퍼스, 레크리에이션 공간) 이용자", False),
        ("· 음환경이 노출 또는 조작 변수(역방향: 이용자 행태·활동이 사운드스케이프 형성·평가에 미치는 영향)", False),
        ("· 관찰 가능한 행태 결과(측정 또는 관찰) · 현장/실험실/VR 실험, 현장 관찰, 자연실험, 센서·빅데이터 관찰 연구", False),
        ("", False),
        ("다 마치면 파일을 저장하고 알려주세요. 원 판정과의 일치도(Cohen's κ)는 rescreen10_agreement.py 로 계산합니다.", False),
    ]
    for i, (text, bold) in enumerate(lines, 1):
        c = ins.cell(row=i, column=1, value=text)
        c.font = Font(bold=bold, size=12 if bold else 11)
    ins.column_dimensions["A"].width = 150

    # ── 판정 시트 ─────────────────────────────────────────────────────────
    ws = wb.create_sheet("Rescreen")
    cols = [("Seq", 6), ("Record ID", 10), ("Year", 7), ("Journal", 24), ("Title", 40), ("Abstract", 90),
            ("DOI", 22), ("Decision", 12), ("Exclusion code", 12), ("Note", 24)]
    for j, (name, width) in enumerate(cols, 1):
        c = ws.cell(row=1, column=j, value=name)
        c.font = Font(bold=True)
        c.fill = head_fill
        c.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        ws.column_dimensions[c.column_letter].width = width
    ws.freeze_panes = "E2"

    dv_dec = DataValidation(type="list", formula1='"Include,Uncertain,Exclude"', allow_blank=True)
    dv_code = DataValidation(type="list", formula1='"' + ",".join(c for c, _, _ in CRITERIA) + '"', allow_blank=True)
    ws.add_data_validation(dv_dec)
    ws.add_data_validation(dv_code)

    for i, no in enumerate(sample, 1):
        r = by_id[no]
        abstract = (r.get("abstract") or "").strip() or "(no abstract in record)"
        vals = [i, no, r["year"], r["journal"], r["title"], abstract, r["doi"], None, None, None]
        row = i + 1
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=row, column=j, value=v)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            if j in (8, 9, 10):
                c.fill = input_fill
        dv_dec.add(f"H{row}")
        dv_code.add(f"I{row}")
        lines_abs = len(abstract) / 95 + 1
        lines_title = len(r["title"]) / 42 + 1
        ws.row_dimensions[row].height = min(409, max(30, 15 * max(lines_abs, lines_title)))

    out = os.path.join(OUTDIR, f"rescreen10_sheet_{stamp}.xlsx")
    wb.save(out)
    print(f"표본 {n}건 / 풀 {len(ids)}건 · 시드 {SEED}")
    print("저장:", out)


if __name__ == "__main__":
    main()
