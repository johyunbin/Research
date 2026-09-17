# -*- coding: utf-8 -*-
"""
Paper32 — 10% 재스크리닝 일치도 (평가자 내 일치, OSF 등록본 Screening reliability)

사용: python rescreen10_agreement.py screening/rescreen10/rescreen10_sheet_<타임코드>.xlsx
원 판정 = screening/screening_results_20260802_221808.csv (저자가 AI 사전분류 전건을 확인·동의한 최종 제목·초록 판정)
- 주지표: 전문 검토로 넘김(Include/Uncertain ↔ INCLUDE/BORDERLINE) vs 배제 — 일치율, Cohen's κ(95% CI), PABAK
- 보조: 3수준(Include/Uncertain/Exclude) κ, 둘 다 배제한 레코드의 배제 코드 일치율
- 불일치 레코드 목록
출력: 같은 폴더 rescreen10_agreement_<타임코드>.md
"""
import csv, math, os, sys
from collections import Counter
from datetime import datetime, timezone, timedelta

from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")
BASE = os.path.dirname(os.path.abspath(__file__))
ORIG = os.path.join(BASE, "screening", "screening_results_20260802_221808.csv")
NEW_MAP = {"include": "INCLUDE", "uncertain": "BORDERLINE", "exclude": "EXCLUDE"}


def kappa(pairs, labels):
    n = len(pairs)
    po = sum(a == b for a, b in pairs) / n
    ca, cb = Counter(a for a, _ in pairs), Counter(b for _, b in pairs)
    pe = sum(ca[l] * cb[l] for l in labels) / n ** 2
    if pe == 1:
        return po, float("nan"), (float("nan"), float("nan"))
    k = (po - pe) / (1 - pe)
    se = math.sqrt(po * (1 - po) / (n * (1 - pe) ** 2))          # Cohen(1960) 근사 표준오차
    return po, k, (k - 1.96 * se, min(1.0, k + 1.96 * se))


def main():
    if len(sys.argv) < 2:
        sys.exit("사용: python rescreen10_agreement.py <채운 xlsx 경로>")
    path = sys.argv[1]
    orig = {int(r["no"]): r for r in csv.DictReader(open(ORIG, encoding="utf-8-sig"))}

    ws = load_workbook(path, data_only=True)["Rescreen"]
    head = [c.value for c in ws[1]]
    col = {name: head.index(name) for name in ("Record ID", "Title", "Decision", "Exclusion code")}
    rows, problems = [], []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[col["Record ID"]] is None:
            continue
        no = int(row[col["Record ID"]])
        dec = str(row[col["Decision"]] or "").strip().lower()
        code = str(row[col["Exclusion code"]] or "").strip().upper()
        if dec not in NEW_MAP:
            problems.append(f"ID {no}: 판정 없음 또는 형식 오류({dec!r})")
            continue
        if dec == "exclude" and not code.startswith("EX"):
            problems.append(f"ID {no}: Exclude 인데 배제 코드 없음")
        rows.append((no, NEW_MAP[dec], code, row[col["Title"]]))
    if problems:
        print("미완성 항목:")
        print("\n".join("  " + p for p in problems))
        sys.exit(1)

    def retain(v):
        return "retain" if v in ("INCLUDE", "BORDERLINE") else "exclude"

    def level3(v):
        return v if v in ("INCLUDE", "BORDERLINE") else "EXCLUDE"

    bin_pairs = [(retain(orig[no]["verdict"]), retain(new)) for no, new, _, _ in rows]
    tri_pairs = [(level3(orig[no]["verdict"]), new) for no, new, _, _ in rows]
    n = len(rows)
    po2, k2, ci2 = kappa(bin_pairs, ["retain", "exclude"])
    po3, k3, ci3 = kappa(tri_pairs, ["INCLUDE", "BORDERLINE", "EXCLUDE"])
    pabak = 2 * po2 - 1
    table = Counter(bin_pairs)
    both_ex = [(orig[no]["verdict"], code) for no, new, code, _ in rows
               if new == "EXCLUDE" and orig[no]["verdict"].startswith("EX")]
    code_agree = sum(a == b for a, b in both_ex)
    disagree = [(no, orig[no]["verdict"], new + (f" ({code})" if code else ""), title)
                for no, new, code, title in rows if retain(orig[no]["verdict"]) != retain(new)]

    stamp = datetime.now(timezone(timedelta(hours=9))).strftime("%Y%m%d_%H%M%S")
    out = os.path.join(os.path.dirname(os.path.abspath(path)), f"rescreen10_agreement_{stamp}.md")
    L = [f"# 10% 재스크리닝 일치도 ({stamp})", "",
         f"- 입력: `{os.path.basename(path)}` · 레코드 {n}건 · 원 판정 `{os.path.basename(ORIG)}`", "",
         "## 주지표: 전문 검토로 넘김 vs 배제", "",
         "| 원 판정 \\ 재판정 | 넘김 | 배제 |", "|---|---|---|",
         f"| 넘김 | {table[('retain', 'retain')]} | {table[('retain', 'exclude')]} |",
         f"| 배제 | {table[('exclude', 'retain')]} | {table[('exclude', 'exclude')]} |", "",
         f"- 일치율 {po2:.1%} · Cohen's κ = {k2:.2f} (95% CI {ci2[0]:.2f} to {ci2[1]:.2f}) · PABAK = {pabak:.2f}", "",
         "## 보조 지표", "",
         f"- 3수준(Include/Uncertain/Exclude) 일치율 {po3:.1%} · κ = {k3:.2f} (95% CI {ci3[0]:.2f} to {ci3[1]:.2f})",
         f"- 둘 다 배제한 {len(both_ex)}건의 배제 코드 일치 {code_agree}건"
         + (f" ({code_agree / len(both_ex):.1%})" if both_ex else ""), "",
         "## 넘김/배제 불일치 레코드", ""]
    if disagree:
        L += ["| ID | 원 판정 | 재판정 | 제목 |", "|---|---|---|---|"]
        L += [f"| {no} | {o} | {nw} | {t} |" for no, o, nw, t in disagree]
    else:
        L.append("없음")
    open(out, "w", encoding="utf-8").write("\n".join(L) + "\n")
    print("\n".join(L))
    print("\n저장:", out)


if __name__ == "__main__":
    main()
