# -*- coding: utf-8 -*-
"""
Paper32 — 원고 수치 ↔ 정본 데이터 대조
집필 중 기억으로 쓴 숫자가 섞이는 것을 막는다. 원고에 등장하는 핵심 수치를 정본에서 재산출해 대조.
사용: python verify_manuscript.py [원고경로]
"""
import math
import sys, os, csv, re
from collections import Counter

sys.stdout.reconfigure(encoding="utf-8")
BASE = os.path.dirname(os.path.abspath(__file__))
FT = os.path.join(BASE, "fulltext")
MA = os.path.join(FT, "ma")
DEFAULT_MS = os.path.join(os.path.dirname(BASE), "01_논문작업",
                          "Manuscript_KO.md")


def rd(p):
    return list(csv.DictReader(open(p, encoding="utf-8-sig")))


def main():
    ms_path = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_MS
    if not os.path.exists(ms_path):
        print(f"⚠️ 원고 없음: {ms_path}"); return 1
    ms = open(ms_path, encoding="utf-8").read()

    cv = rd(os.path.join(FT, "corpus_v4_verdicts.csv"))
    t1 = rd(os.path.join(FT, "table1_v2.csv"))
    q = rd(os.path.join(FT, "quality_v2.csv"))
    ec = rd(os.path.join(FT, "evidence_counts_v2.csv"))

    inc = [r for r in t1 if r["verdict"] == "FINAL_INCLUDE"]
    facts = {}

    facts["최종 포함 96"] = sum(1 for r in cv if r["final_verdict"] == "FINAL_INCLUDE")
    facts["민감도 4"] = sum(1 for r in cv if r["final_verdict"] == "SENS_ONLY")
    facts["DB 갈래 81"] = sum(1 for r in cv if r["final_verdict"] == "FINAL_INCLUDE"
                             and r["source"] == "db-search")
    facts["인용추적 15"] = sum(1 for r in cv if r["final_verdict"] == "FINAL_INCLUDE"
                              and r["source"] == "citation-tracking")

    ys = [int(r["year"]) for r in inc if r["year"].isdigit()]
    facts["2020년 이후 68"] = sum(1 for y in ys if y >= 2020)
    facts["최소 연도 1975"] = min(ys)
    facts["중앙 연도 2022"] = sorted(ys)[len(ys) // 2]

    st = Counter(r["setting"] for r in inc)
    facts["street 36"] = st["street"]
    facts["park 28"] = st["park"]
    de = Counter(r["design"] for r in inc)
    facts["field experiment 17"] = de["field experiment"]

    qt = Counter(r["quality"] for r in inc)
    facts["high 22"] = qt["high"]
    facts["moderate 40"] = qt["moderate"]
    facts["low 34"] = qt["low"]

    # 국가
    import importlib.util
    spec = importlib.util.spec_from_file_location("g7", os.path.join(BASE, "make_fig7_geo_time.py"))
    g7 = importlib.util.module_from_spec(spec)
    try:
        spec.loader.exec_module(g7)
    except SystemExit:
        pass
    ext = rd(os.path.join(FT, "corpus_v4_extraction.csv"))
    cc = Counter()
    for r in ext:
        if r["uid"] not in {x["uid"] for x in inc}:
            continue
        for c in g7.norm_countries(r["country"]):
            cc[c] += 1
    facts["China 40"] = cc["China"]

    dirn = Counter()
    for r in ec:
        if r["kind"] == "direction_x_domain":
            dirn[r["col"]] += int(r["n"])
    facts["forward 110"] = dirn["forward"]
    facts["reverse 74"] = dirn["reverse"]
    facts["both 22"] = dirn["both"]

    gen = Counter()
    for r in ec:
        if r["kind"] == "generation_x_band":
            gen[(r["row"], r["col"])] = int(r["n"])
    facts["G3 2010s 3"] = gen[("2010–2019", "G3")]
    facts["G3 2020s 21"] = gen[("2020–", "G3")]
    facts["G1 2020s 38"] = gen[("2020–", "G1")]
    facts["G2 2020s 28"] = gen[("2020–", "G2")]

    multi = sum(1 for r in inc if len([g for g in r["measure_gen"].split(";")
                                       if g.strip() in ("G1", "G2", "G3")]) >= 2)
    facts["2세대 이상 17"] = multi

    air = sum(int(r["n"]) for r in ec
              if r["kind"] == "domain_x_source" and r["col"] == "항공기소음")
    facts["항공기 4"] = air

    # ── 대조 ───────────────────────────────────────────────────────
    # ⚠️ 라벨에 기대값을 박아두면 정본이 갱신될 때 라벨이 낡는다(실제로 그랬다).
    #    데이터에서 산출한 값이 **원고 본문에 실제로 등장하는지**를 본다.
    ok, bad = [], []
    for label, val in facts.items():
        name = re.sub(r"\s*\d+$", "", label)
        # 숫자가 원고 어딘가에 있는가(천단위 콤마 허용).
        # ⚠️ 한글본에서는 "113건"처럼 숫자 뒤에 한글이 붙어 `\b`가 성립하지 않는다.
        #    앞뒤 경계를 "숫자가 아닌 것"으로 완화한다.
        pats = [rf"(?<!\d){val}(?!\d)", rf"(?<!\d){val:,}(?!\d)"]
        found = any(re.search(p, ms) for p in pats)
        (ok if found else bad).append((name, val))

    must = {"코퍼스 98": r"\b98\b", "DB 갈래 81": r"\b81\b", "인용추적 15": r"\b15\b",
            "보조검색 2": r"보조 색인 경로에서는[^\n]*2편을 포함했다", "현장조작 비율 19%": r"19%", "역방향 비율 40%": r"40%"}
    missing = [k for k, pat in must.items() if not re.search(pat, ms)]

    # 메타분석 수치가 원고와 일치하는가
    # ⚠️ 기대값을 여기 박아두면 정본이 바뀔 때 검증기만 낡는다(실제로 낡았다).
    #    ma_forest_data.json(= ma_v2.py 산출)에서 읽어 원고에 그 값이 있는지 본다.
    import json
    fj = json.load(open(os.path.join(MA, "ma_forest_data.json"), encoding="utf-8"))
    NAME = {"walking": "MA1 walking", "staying": "MA2 staying",
            "social": "MA3 social", "correlation": "MA4 correlation"}
    ma_bad = []
    # ★ 2026-09-17 표기 변경: 효과크기·신뢰구간은 소수 둘째 자리, CI 는 그림·표와 같은 [하한, 상한],
    #   p 는 소수 셋째 자리(선행 0 포함). 본문 문장에 "95% CI [..]" 가 그대로 있는지 본다.
    def f2(x):
        return f"{x:+.2f}".replace("-", "−")

    for key, nm in NAME.items():
        p = fj[key]["pooled"]
        est, lo, hi = p["est"], p["lo"], p["hi"]
        if fj[key].get("back_r"):
            est, lo, hi = math.tanh(est), math.tanh(lo), math.tanh(hi)
        want = f"= {f2(est)}, 95% CI [{f2(lo)}, {f2(hi)}]"
        if want not in ms:
            ma_bad.append(f"{nm} '{want}' — 원고 본문에 없음")
        if f"*p* = {p['p']:.3f}" not in ms:
            ma_bad.append(f"{nm} p = {p['p']:.3f} — 원고 본문에 없음")
        if not re.search(rf"\bk\b[^\n]{{0,12}}=\s*{p['k']}\b|\|\s*{p['k']}\s*\|", ms):
            ma_bad.append(f"{nm} k={p['k']} — 원고에 없음")

    # leave-one-out: Table 3 의 범위 행과 3.5절 문장을 반올림 전 값(ma_sensitivity_v2_raw.json)과 대조
    # (D7 재계산 뒤 LOO 수치가 구값으로 남아 있던 것을 게이트가 적발한 이력 — 2026-09-16)
    import json as _json
    raw = _json.load(open(os.path.join(MA, "ma_sensitivity_v2_raw.json"), encoding="utf-8"))
    for pre in ("MA1", "MA2", "MA3", "MA4"):
        loo = [r for r in raw if r["cluster"].startswith(pre) and r["analysis"].startswith("LOO")]
        es = [r["r"] if "r" in r else r["est"] for r in loo]
        ps = [r["p"] for r in loo]
        row = f"{f2(min(es))} to {f2(max(es))} | — | {min(ps):.3f}–{max(ps):.3f}"   # p 범위는 en dash(2026-09-17)
        if row not in ms:
            ma_bad.append(f"{pre} LOO 범위 '{row}' — Table 3 에 없음")
    m3 = sorted([r for r in raw if r["cluster"].startswith("MA3") and r["analysis"].startswith("LOO")],
                key=lambda r: r["est"])
    low = f"*g* = {f2(m3[0]['est'])}, *p* = {m3[0]['p']:.3f}"
    if low not in ms:
        ma_bad.append(f"MA3 LOO 최저 '{low}' — 3.5절에 없음")

    # 원고의 Table 2·3 이 생성본(build_ma_char_table.py · build_ma_sensitivity_table.py)과 같은가
    for name, fn in (("Table 2", "ma_char_table.md"), ("Table 3", "ma_sensitivity_table.md")):
        gen = open(os.path.join(FT, fn), encoding="utf-8").read().strip()
        if gen not in ms:
            ma_bad.append(f"{name} 이 생성본 fulltext/{fn} 과 다르다 — 스크립트를 다시 돌려 붙여 넣을 것")

    # 사용자가 다시 그린 Fig. 1(01_논문작업/Figure.pptx)의 흐름도 숫자 = PRISMA 정본 수치인가
    # (2026-09-17 그림 1·8 을 사용자 작도본으로 교체 — 스크립트가 그리지 않으므로 수치가 바뀌면 그림이 낡는다)
    pptx_bad = []
    pptx_path = os.path.join(os.path.dirname(BASE), "01_논문작업", "Figure.pptx")
    try:
        from pptx import Presentation
        spec = importlib.util.spec_from_file_location("f1", os.path.join(BASE, "make_fig1_prisma_v2.py"))
        f1 = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(f1)
        DB, O = f1.DB, f1.O
        want = {DB[k] for k in ("identified", "wos", "scopus", "pubmed", "duplicates", "screened", "excluded",
                                "sought", "not_retrieved", "assessed", "ft_excluded")}
        want |= {O[k] for k in ("cit", "supp", "dup", "auto", "screened", "excluded", "sought", "not_retrieved",
                                "assessed", "ft_excluded")}
        want |= {O["cit"] + O["supp"], O["dup"] + O["auto"], f1.D["n_included"], f1.D["n_sens"]}
        allowed = want | {v for _, v in DB["ftx"]} | {v for _, v in O["ftx"]}

        def shapes(ss):
            for s in ss:
                if s.shape_type == 6:
                    yield from shapes(s.shapes)
                else:
                    yield s
        slide = next(sl for sl in Presentation(pptx_path).slides
                     if any(s.has_text_frame and "Identification of studies via databases" in s.text_frame.text
                            for s in shapes(sl.shapes)))
        txt = " ".join(s.text_frame.text for s in shapes(slide.shapes) if s.has_text_frame)
        got = {int(x.replace(",", "")) for x in re.findall(r"n = ([\d,]+)", txt)}
        if want - got:
            pptx_bad.append(f"Figure.pptx 흐름도에 없는 정본 수치: {sorted(want - got)}")
        if got - allowed:
            pptx_bad.append(f"Figure.pptx 흐름도에만 있는 수치(정본에 없음): {sorted(got - allowed)}")
    except FileNotFoundError:
        pptx_bad.append("01_논문작업/Figure.pptx 없음")
    ma_bad += pptx_bad

    print(f"=== 원고 수치 검증 — {os.path.basename(ms_path)} ===\n")
    print(f"데이터 산출값이 원고에 존재: {len(ok)} / 누락 {len(bad)}")
    for name, v in bad:
        print(f"  ❌ {name} = {v} → 원고에서 못 찾음")
    if not bad:
        print("  ✅ 전 항목 일치")
    print(f"\n원고 필수 수치 누락: {missing if missing else '없음'}")
    print(f"메타분석 수치 대조: {ma_bad if ma_bad else '✅ 일치'}")

    # 원고에 남은 플레이스홀더
    todo = re.findall(r"\*\[.*?\]\*", ms)
    print(f"\n미작성 구간 {len(todo)}개: {[t[:40] for t in todo]}")
    return 1 if (bad or ma_bad) else 0


if __name__ == "__main__":
    sys.exit(main())
